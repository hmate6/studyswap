from uuid import uuid4
from time import time
import asyncio
import threading
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement

def makeDocument(elements, token):
    doc = Document()

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.3)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    doc.add_heading('Automated Outline', level=1)

    for element in elements:
        para = doc.add_paragraph(f"{element}", style='ListBullet')
        for run in para.runs:
            run.font.name = 'Calibri'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri') 
            run.font.size = Pt(12)

    doc.save(f"./static/files/{token}.docx")
    return f"{token}.docx"

def gen_token():
    tim = str(int(time()))
    tim = tim[len(tim)-4::]

    return str(uuid4()) + tim
